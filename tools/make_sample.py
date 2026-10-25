"""Builds the sample test file in the judges' format, plus sample-data.js for the
"Load sample data" button. Run: python tools/make_sample.py
The rows are deliberately messy: contested pins, an existing-store pin, a city/pin
contradiction, a duplicate, blank investment, text in number columns, etc.
"""
import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

APP_HEADERS = ["Application ID", "Applicant name", "City", "State", "Pin code", "Investment capacity (₹ lakh)",
               "Retail experience (years)", "Current / past business", "Space available", "Space size (sq ft)",
               "Floor", "Target launch (months)", "Preferred language", "Notes from applicant"]
K = "Karnataka"
APPS = [
    ["APP-001", "Ravi Kumar", "Mysuru", K, 570001, 30, 4, "Runs a footwear store", "Rented", 450, "Ground", 4, "Kannada", "Wants to open before Diwali"],
    ["APP-002", "Priya Shetty", "Mysuru", K, 570001, 40, 6, "Managed a fashion retail outlet in a mall", "Owned", 600, "Ground", 3, "English", ""],
    ["APP-003", "Mohammed Irfan", "Bengaluru", K, 560034, 45, 8, "Owns an optical shop", "Owned", 500, "Ground", 3, "Hindi", "Shop is on the main road in Koramangala"],
    ["APP-004", "Sneha Rao", "Bengaluru", K, 560008, 28, 2, "IT professional; wife runs a boutique", "Rented", 350, "Ground", 3, "English", "Bank loan for the remaining amount"],
    ["APP-005", "Arjun Patil", "Belagavi", K, 590006, 12, 1, "Mobile repair shop", "Not yet identified", "", "Not known", 2, "Kannada", ""],
    ["APP-006", "Lakshmi Narayan", "Tumakuru", K, 572101, 22, 0, "Retired government school teacher", "Not yet identified", 400, "Ground", 6, "Kannada", "Son will help run the store"],
    ["APP-007", "Vikram Hegde", "Mangaluru", K, 560010, 32, 5, "Runs a stationery shop", "Rented", 400, "Ground", 3, "English", ""],
    ["APP-008", "Fatima Begum", "Hubballi-Dharwad", K, 580021, 35, 5, "Owns a garment shop", "Owned", 800, "Upper", 3, "Hindi", "First floor above my garment shop"],
    ["APP-009", "Suresh Gowda", "Hubballi-Dharwad", K, 580029, 30, 3, "Kirana store owner", "Rented", 500, "Ground", 4, "Kannada", ""],
    ["APP-010", "Anitha Kumari", "Mysuru", K, 570012, 28, 3, "Runs a bakery", "Rented", 420, "Ground", 4, "Kannada", ""],
    ["APP-011", "Rahul Jain", "Bengaluru", K, 560067, 50, 10, "Runs two franchise outlets of a bakery chain", "Rented", 700, "Mall unit", 5, "Hindi", ""],
    ["APP-012", "Deepa Nair", "Bengaluru", K, 560070, 40, 3, "Runs a beauty salon", "Not yet identified", "", "Not known", 4, "English", "Looking for a shop on the main road"],
    ["APP-013", "Kiran Desai", "Davanagere", K, 577002, 20, 2, "Agricultural products dealer", "Rented", 380, "Ground", 3, "Kannada", ""],
    ["APP-014", "Meera Joshi", "Davanagere", K, 577004, 25, 6, "Ran a pharmacy for 6 years", "Owned", 500, "Ground", 4, "Hindi", ""],
    ["APP-015", "Naveen Reddy", "Hubballi-Dharwad", K, 580030, "", 3, "Electronics dealer", "Rented", 450, "Ground", 3, "English", "Will confirm budget after talking to family"],
    ["APP-016", "Ravi Kumar", "Mysuru", K, 570002, 25, 4, "Footwear store", "Not yet identified", "", "Not known", 6, "Kannada", "Applying again in case first form did not go through"],
    ["APP-017", "Sanjay Kulkarni", "Belagavi", K, 590001, 30, 4, "Runs a sports goods store", "Owned", 550, "Ground", 3, "Kannada", ""],
    ["APP-018", "Pooja Hegde", "Tumakuru", K, 572102, 14, 0, "Homemaker", "Rented", 250, "Ground", 2, "Kannada", "Just exploring options"],
    ["APP-019", "Abdul Rahman", "Mangaluru", K, 575002, 30, 7, "Electronics showroom owner", "Owned", 1500, "Ground", 12, "English", "Also selling another eyewear brand in the showroom"],
    ["APP-020", "Harish Shetty", "Mangaluru", K, 575003, 26, 4, "Distributor of FMCG goods", "Rented", 420, "Mall unit", 4, "English", ""],
    ["APP-021", "Gita Menon", "Bengaluru", K, 560043, 60, 12, "Former store manager at a large retail chain", "Not yet identified", "", "Not known", 8, "English", "Will quit my job and run it full-time"],
    ["APP-022", "Ramesh Babu", "Mysuru", K, 570008, "35 lakhs", "seven", "Ran a mobile phone shop", "Rented", "600 sqft", "ground floor", 2, "Hindi", "Partnering with my brother"],
    ["APP-023", "Shalini Prakash", "Tumakuru", K, 572103, 24, 3, "Runs a gift and stationery store", "Owned", "", "Not known", 5, "Kannada", ""],
]
TER_HEADERS = ["Territory ID", "City", "State", "Pin codes covered", "City tier", "Open slots", "Priority"]
TERS = [
    ["TER-01", "Bengaluru", K, "560008, 560010, 560011, 560034, 560043", "Tier 1", 2, "High"],
    ["TER-02", "Bengaluru", K, "560066, 560067, 560070, 560076, 560078", "Tier 1", 1, "Medium"],
    ["TER-03", "Mysuru", K, "570001, 570002, 570008, 570009, 570010", "Tier 2", 2, "High"],
    ["TER-04", "Mangaluru", K, "575001, 575002, 575003", "Tier 2", 1, "Medium"],
    ["TER-05", "Hubballi-Dharwad", K, "580020, 580021, 580029, 580030", "Tier 2", 2, "High"],
    ["TER-06", "Belagavi", K, "590001, 590006", "Tier 2", 1, "Low"],
    ["TER-07", "Tumakuru", K, "572101, 572102, 572103", "Tier 3", 2, "High"],
    ["TER-08", "Davanagere", K, "577001, 577002, 577004", "Tier 3", 1, "Medium"],
]
STORE_HEADERS = ["Store ID", "City", "Pin code", "Opened (year)"]
STORES = [["ST-01", "Bengaluru", 560001, 2019], ["ST-02", "Bengaluru", 560034, 2020], ["ST-03", "Bengaluru", 560038, 2021],
          ["ST-04", "Bengaluru", 560066, 2022], ["ST-05", "Bengaluru", 560076, 2023], ["ST-06", "Mysuru", 570009, 2021],
          ["ST-07", "Mysuru", 570017, 2024], ["ST-08", "Mangaluru", 575001, 2022], ["ST-09", "Hubballi-Dharwad", 580020, 2023],
          ["ST-10", "Belagavi", 590001, 2024]]
GLOSSARY = [
    ["Column", "Meaning"],
    ["Investment capacity (₹ lakh)", "Money the applicant can put into the store; 1 lakh = ₹1,00,000"],
    ["Retail experience (years)", "Years spent running or managing any shop; 0 if none"],
    ["Space available", "Owned / Rented / Not yet identified"],
    ["Floor", "Ground / Upper / Mall unit / Not known"],
    ["Target launch (months)", "How many months from now the applicant wants to open"],
    ["Pin codes covered", "6-digit pin codes inside the territory, comma-separated"],
    ["Open slots", "How many more stores the brand wants to open in the territory"],
    ["Priority", "How keen the brand is to grow here right now: High / Medium / Low"],
]


def sheet(wb, title, headers, rows, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, h in enumerate(headers):
        width = max(len(str(h)), *(len(str(r[i])) for r in rows)) if rows else len(str(h))
        ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = min(45, max(10, width + 2))
    ws.freeze_panes = "A2"


def main():
    os.makedirs(os.path.join(ROOT, "sample"), exist_ok=True)
    wb = Workbook()
    sheet(wb, "Applications", APP_HEADERS, APPS, first=True)
    sheet(wb, "Territories", TER_HEADERS, TERS)
    sheet(wb, "Existing stores", STORE_HEADERS, STORES)
    sheet(wb, "Glossary", GLOSSARY[0], GLOSSARY[1:])
    out = os.path.join(ROOT, "sample", "FranchiseFlow_Sample_Test_Data.xlsx")
    wb.save(out)

    blank = Workbook()
    sheet(blank, "Applications", APP_HEADERS, [], first=True)
    sheet(blank, "Territories", TER_HEADERS, [])
    sheet(blank, "Existing stores", STORE_HEADERS, STORES)
    sheet(blank, "Glossary", GLOSSARY[0], GLOSSARY[1:])
    blank.save(os.path.join(ROOT, "sample", "FranchiseFlow_Blank_Template.xlsx"))

    sheets = [
        {"name": "Applications", "rows": [APP_HEADERS] + APPS},
        {"name": "Territories", "rows": [TER_HEADERS] + TERS},
        {"name": "Existing stores", "rows": [STORE_HEADERS] + STORES},
    ]
    with open(os.path.join(ROOT, "sample-data.js"), "w", encoding="utf-8") as f:
        f.write("/* Generated by tools/make_sample.py - same rows as sample/FranchiseFlow_Sample_Test_Data.xlsx */\n")
        f.write("window.SAMPLE_SHEETS = " + json.dumps(sheets, ensure_ascii=False, indent=1) + ";\n")
    print("wrote", out)


if __name__ == "__main__":
    main()
